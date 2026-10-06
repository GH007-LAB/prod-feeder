# -*- coding: utf-8 -*-
"""
007 Metals - เติม Excel "รายงานขายประจำวัน" อัตโนมัติจากแอปปิดรอบ 007 + Express

เฟสนี้ (ทดลองคู่ขนาน): **ไม่แตะชีตของพนักงาน** — เขียนลง "ชีตคู่แฝด" ชื่อ "<วัน>_auto"
(เช่น ชีต "6" → "6_auto") สร้างด้วย copy_worksheet ของชีตวันจริง (ไม่มีก็ชีตวันก่อนหน้าล่าสุด)
แล้วล้างเฉพาะช่องค่าที่สคริปต์รับผิดชอบก่อนเติม · ช่องสูตรห้ามแตะ (ตรวจซ้ำทุกเซลล์ก่อนเขียน)
รันซ้ำได้ (idempotent): ชีต _auto เดิมถูกลบแล้วสร้างใหม่

แหล่งข้อมูล
  · Express DBF ของสาขา (cfg_{BR}.txt → SRC / PROXY_URL เหมือน cash_feed)
      ARTRN  RECTYP 3=IV 0=AI 5=SR 9=RE (DOCSTAT 'C' = ยกเลิก ข้าม)
      ARRCPIT RCPNUM(RE) → DOCNUM(ใบที่ชำระ) + RCVAMT · ARMAS ชื่อ/ที่อยู่ · istab (TABTYP 40) เขตการขาย
  · Supabase RPC cash_round_detail(p_branch, p_date) ด้วย service key (docs/expenses/round/traffic)
    + ตาราง employees (full_name ผู้ส่งรอบ) — ฟิลด์ qr / noncash_channel / traffic (sql/005)
    ยังไม่มีก็ทำงานได้: ไม่มี noncash_channel = ถือเป็นโอน (transfer) · ไม่มี traffic = เว้นว่าง

ไฟล์เป้าหมาย (Drive Mirror ของ 007skn0777): ดู excel_path() · override โฟลเดอร์รากด้วย
env CASH_EXCEL_ROOT=<โฟลเดอร์ "ไดรฟ์ของฉัน (...)"> หรือระบุไฟล์ตรง ๆ ด้วย --file

การเขียนไฟล์จริง (--write):
  ไฟล์ถูกเปิดค้าง (~$ lock ของ Excel) → ข้าม + log · สำรอง <file>.bak_YYMMDD_HHMM (เก็บ 7 ล่าสุด)
  → โหลด → เติม → save ไฟล์ชั่วคราวในโฟลเดอร์เดียวกัน → ตรวจว่าไฟล์ต้นทางไม่ถูกแก้ระหว่างนั้น
  → os.replace → เปิดอ่านซ้ำ ยืนยันจำนวนชีต + เซลล์ตัวอย่าง
  --out <path> = เขียนผลลงสำเนาที่ path แทน (ไม่สำรอง ไม่แตะไฟล์ต้นทาง) ใช้ทดสอบ

usage:
  python3 cash_excel_fill.py BR YYYY-MM-DD [--dry-run|--write] [--out path] [--file src.xlsx]
                             [--show]
     ค่าเริ่ม = --dry-run (อ่าน + สรุป ไม่เขียนอะไร) · --show พิมพ์ทุกเซลล์ที่จะเติม (stdout)
  cash_poller.py เรียก fill_after_export() หลัง cash_mark_exported เมื่อ CASH_EXCEL_ENABLED=1

log: "CASH_EXCEL: ..." · "ERROR: CASH_EXCEL ..." เฉพาะพังจริง (alert_scan จับ)
ไม่ลงชื่อลูกค้าใน log (PDPA) — log เฉพาะจำนวน/ยอดรวม
"""
import sys, os, re, glob, time, json, shutil, datetime, urllib.request, urllib.error, urllib.parse

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import so_push as S
import cash_feed as F

try:
    import openpyxl
    from openpyxl.comments import Comment
except ImportError:  # pragma: no cover
    openpyxl = None

DRIVE_GLOB = "ไดรฟ์ของฉัน (007skn0777@gmail.com)"
AUTO_SUFFIX = "_auto"
BACKUP_KEEP = 7
DEFAULT_FLOAT = 5000.0
ONLINE_AREA = "เพจ007"          # ที่ SKN ใช้กับ Shopee/Lazada/TikTok (BK/PPS ยังไม่เจอตัวอย่าง)
COMMENT_AUTHOR = "cash_excel_fill"
NOTE_NO_CHANNEL = "ยังไม่เลือกช่องทางในแอป"
BANK_LABEL = {"transfer": "KSME", "qr": "KSHOP"}
CHANNEL_COL = {"cash": "F", "credit": "G", "transfer": "H", "qr": "I"}
TH_WEEKDAY = ["จันทร์", "อังคาร", "พุธ", "พฤหัสบดี", "ศุกร์", "เสาร์", "อาทิตย์"]

# ---- โครงชีต (FORM_LAYOUT.md — เหมือนกันทั้ง 3 สาขา) ----
IV_ROWS = range(3, 63)        # 60 แถว
AI_ROWS = range(65, 73)       # 8
SR_ROWS = range(75, 82)       # 7
CHQ_ROWS = range(85, 87)      # เช็ค — ไม่มีข้อมูลในแอป ล้างทิ้งในชีต _auto
PAID_ROWS = range(90, 110)    # 20 วางบิลชำระแล้ว
EXP_ROWS = range(114, 130)    # 16 ค่าใช้จ่าย
KPI_CELLS = {"visitors": "N3", "closed": "N4", "lost_brand": "N8", "lost_price": "N9",
             "lost_no_stock": "N10", "lost_survey": "N11", "lost_qty": "N12",
             "lost_delivery": "N13", "lost_other": "N14"}
KPI_NOTE_CELL = "N16"
SINGLE_CELLS = ("C123", "C125", "C130", "C131")


def clear_targets():
    """ทุกเซลล์ค่าที่ชีต _auto รับผิดชอบ (ล้างก่อนเติม) — เซลล์สูตรถูกกรองทิ้งอีกชั้นตอนล้าง"""
    cells = []
    for rows, cols in ((IV_ROWS, "ABCDEFGHIL"), (AI_ROWS, "ABCDEFGHIL"), (SR_ROWS, "ABCDEFGHIL"),
                       (CHQ_ROWS, "ABCDFGHIK"), (PAID_ROWS, "ABCDEFGHIJL"), (EXP_ROWS, "GK")):
        for r in rows:
            cells += ["%s%d" % (c, r) for c in cols]
    cells += list(SINGLE_CELLS) + list(KPI_CELLS.values()) + [KPI_NOTE_CELL]
    return cells


# =====================================================================================
# ไฟล์ / ชีต
# =====================================================================================
def drive_root():
    env = os.environ.get("CASH_EXCEL_ROOT")
    if env:
        return env
    home = os.path.expanduser("~")
    p = os.path.join(home, DRIVE_GLOB)
    if os.path.isdir(p):
        return p
    for q in sorted(glob.glob(os.path.join(home, "ไดรฟ์ของฉัน*"))):
        if os.path.isdir(os.path.join(q, "รายงานขายสาขาสกลนคร")):
            return q
    return p


def excel_path(branch, day, root=None):
    """ไฟล์รายเดือนของสาขา · MM = เดือน 2 หลัก, yy = พ.ศ. 2 หลัก, yyyy = พ.ศ. 4 หลัก"""
    be = day.year + 543
    yy, yyyy, mm = "%02d" % (be % 100), str(be), "%02d" % day.month
    rel = {
        "SKN": ("รายงานขายสาขาสกลนคร", "รายงานขายSKNปี" + yy, "รายงานขายสกลนคร%s%s.xlsx" % (yy, mm)),
        "BK": ("รายงานขายสาขาบึงกาฬ", "รายงานขายBK", "รายงานขายปีBK" + yyyy,
               "รายงานขายบึงกาฬ%s%s.xlsx" % (yy, mm)),
        "PPS": ("รายงานขายสาขาโพนพิสัย", "รายงานขายPPSปี" + yy, "รายงานขายโพนพิสัย%s%s.xlsx" % (yy, mm)),
    }.get(branch)
    if not rel:
        raise ValueError("ไม่รู้จักสาขา %s" % branch)
    return os.path.join(root or drive_root(), *rel)


_DAY_RE = re.compile(r"^\s*(\d{1,2})\s*(\(?\s*หยุด\s*\)?)?\s*$")


def day_sheet_map(wb):
    """{เลขวัน: ชื่อชีต} — รองรับ '4(หยุด)' / '4 หยุด ' / '4หยุด' · ข้ามชีต _auto"""
    out = {}
    for name in wb.sheetnames:
        if name.endswith(AUTO_SUFFIX):
            continue
        m = _DAY_RE.match(name)
        if m:
            out.setdefault(int(m.group(1)), name)
    return out


def template_sheet(wb, day_no):
    """ชีตวันจริง ถ้าไม่มี (พนักงานยังไม่สร้าง) ใช้วันก่อนหน้าล่าสุด → (ชื่อ, same_day)"""
    m = day_sheet_map(wb)
    if day_no in m:
        return m[day_no], True
    prev = [d for d in m if d < day_no]
    if prev:
        return m[max(prev)], False
    if m:
        return m[min(m)], False
    return None, False


def excel_locked(path):
    """Excel/Numbers เปิดค้าง → มีไฟล์ ~$<ชื่อ> ข้าง ๆ (Drive Mirror ซิงก์มาด้วย)"""
    d, b = os.path.split(path)
    for cand in ("~$" + b, "~$" + b[2:], ".~lock.%s#" % b):
        if os.path.exists(os.path.join(d, cand)):
            return cand
    return None


# =====================================================================================
# Express
# =====================================================================================
def _find(src, name):
    """ชื่อไฟล์ DBF แบบไม่สนตัวพิมพ์ (istab.DBF ตัวเล็ก) — โหมด proxy คืนชื่อมาตรฐาน"""
    p = os.path.join(src, name)
    if S.PROXY or os.path.exists(p):
        return p
    try:
        for f in os.listdir(src):
            if f.lower() == name.lower():
                return os.path.join(src, f)
    except OSError:
        pass
    return p


def _num(v):
    return round(float(v or 0), 2)


_IV_RE = re.compile(r"IV\d{5,}")


def load_express(src, day, extra_re=()):
    """อ่าน DBF → dict ที่ build_plan ใช้ (ไม่มีการเรียกเครือข่าย)
    extra_re = เลข RE จากแอป (เช่นใบยกมา) ที่ต้องดึงรายการ ARRCPIT ด้วย"""
    artrn = _find(src, "ARTRN.DBF")
    fields = {"RECTYP", "DOCNUM", "DOCDAT", "CUSCOD", "AREACOD", "ADVNUM", "ADVAMT", "NETAMT",
              "DOCSTAT", "YOUREF", "SONUM"}
    ivs, ais, srs, res = [], [], [], {}
    want_re = set(extra_re)
    for r in S.read_dbf(artrn, fields=fields):
        t = (r.get("RECTYP") or "").strip()
        doc = (r.get("DOCNUM") or "").strip()
        if not doc or (r.get("DOCSTAT") or "").strip() == F.CANCELLED:
            continue
        if r.get("DOCDAT") == day:
            if t == "3":
                ivs.append(r)
            elif t == "0":
                ais.append(r)
            elif t == "5":
                srs.append(r)
            elif t == "9":
                res[doc] = r
        if t == "9" and doc in want_re and doc not in res:
            res[doc] = r
    # รายการชำระของ RE
    lines = {}
    rc = _find(src, "ARRCPIT.DBF")
    if S._file_exists(rc):
        need = set(res) | want_re
        for r in S.read_dbf(rc, fields={"RCPNUM", "DOCNUM", "RECTYP", "RCVAMT"}):
            rcp = (r.get("RCPNUM") or "").strip()
            if rcp in need:
                lines.setdefault(rcp, []).append({
                    "doc_no": (r.get("DOCNUM") or "").strip(),
                    "rectyp": (r.get("RECTYP") or "").strip(),
                    "amount": _num(r.get("RCVAMT"))})
    # IV ที่ RE ชำระ (รวมใบวันก่อน ๆ)
    day_iv = {(r.get("DOCNUM") or "").strip() for r in ivs}
    need_iv = {ln["doc_no"] for ls in lines.values() for ln in ls if ln["rectyp"] == "3"} - day_iv
    iv_index = {}
    for r in ivs:
        iv_index[(r.get("DOCNUM") or "").strip()] = r
    if need_iv:
        for r in S.read_dbf(artrn, fields=fields):
            doc = (r.get("DOCNUM") or "").strip()
            if doc in need_iv and (r.get("RECTYP") or "").strip() == "3":
                iv_index[doc] = r
    # ลูกค้า: ชื่อ + ที่อยู่ (หาพื้นที่)
    cus_need = {(r.get("CUSCOD") or "").strip()
                for r in list(ivs) + ais + srs + list(res.values()) + list(iv_index.values())}
    cus_need.discard("")
    customers = {}
    if cus_need:
        for r in S.read_dbf(_find(src, "ARMAS.DBF"),
                            fields={"CUSCOD", "CUSNAM", "ADDR01", "ADDR02", "ADDR03", "AREACOD"}):
            c = (r.get("CUSCOD") or "").strip()
            if c in cus_need:
                customers[c] = {"name": F.norm(r.get("CUSNAM")),
                                "addr": " ".join(F.norm(r.get(k)) for k in ("ADDR01", "ADDR02", "ADDR03")),
                                "area": (r.get("AREACOD") or "").strip()}
    areas = {}
    ist = _find(src, "istab.DBF")
    if S._file_exists(ist):
        for r in S.read_dbf(ist, fields={"TABTYP", "TYPCOD", "SHORTNAM", "TYPDES"}):
            if (r.get("TABTYP") or "").strip() == "40":
                areas[(r.get("TYPCOD") or "").strip()] = F.norm(r.get("SHORTNAM") or r.get("TYPDES"))
    key = lambda r: (r.get("DOCNUM") or "")
    return {"day": day, "ivs": sorted(ivs, key=key), "ais": sorted(ais, key=key),
            "srs": sorted(srs, key=key), "res": res, "lines": lines, "iv_index": iv_index,
            "customers": customers, "areas": areas, "has_arrcpit": S._file_exists(rc)}


_AMPHOE = re.compile(r"(?:อ\.|อำเภอ)\s*([^\s\d]+)")
_PROVINCE = re.compile(r"(?:จ\.|จังหวัด)\s*([^\s\d]+)")
ONLINE_NAMES = {"SHOPEE", "LAZADA", "TIKTOK", "SHOPIFY"}


def area_label(r, ex):
    """พื้นที่ช่อง C — พนักงานเขียนชื่ออำเภอ (ตามที่อยู่ลูกค้า) / 'เพจ007' ถ้าขายออนไลน์
    ลำดับ: ออนไลน์ → อำเภอจาก ARMAS → จังหวัด → ชื่อเขตใน istab (ยกเว้น Walk In) → รหัสดิบ"""
    cus = (r.get("CUSCOD") or "").strip()
    code = (r.get("AREACOD") or "").strip()
    c = ex["customers"].get(cus) or {}
    if (code.endswith("อล") or code == "อน" or cus.startswith("O")
            or (c.get("name") or "").upper() in ONLINE_NAMES):
        return ONLINE_AREA
    addr = c.get("addr") or ""
    m, p = _AMPHOE.search(addr), _PROVINCE.search(addr)
    if m:
        a = m.group(1)
        if a.startswith("เมือง"):           # อ.เมือง / อ.เมืองสกลนคร → ชื่อจังหวัด (แบบที่พนักงานเขียน)
            a = a[len("เมือง"):] or (p.group(1) if p else a)
        return a
    if p:
        return p.group(1)
    name = ex["areas"].get(code) or ex["areas"].get(c.get("area") or "")
    if name and name.lower() != "walk in":
        return name
    return code or None


def cus_name(r, ex):
    cus = (r.get("CUSCOD") or "").strip()
    return (ex["customers"].get(cus) or {}).get("name") or cus or None


# =====================================================================================
# แอปปิดรอบ (Supabase)
# =====================================================================================
def fetch_round(cfg, branch, day):
    """cash_round_detail(p_branch, p_date) → dict (round อาจเป็น None) · RPC ยังไม่มี → None"""
    try:
        return F.RPC(cfg, "cash_round_detail", {"p_branch": branch, "p_date": day.isoformat()}) or {}
    except F.RpcMissing:
        S.log("CASH_EXCEL: %s ยังไม่มี RPC cash_round_detail — เติมจาก Express อย่างเดียว" % branch)
        return None


def _rest_get(cfg, path, timeout=15):
    url = cfg["SUPABASE_URL"].rstrip("/") + path
    req = urllib.request.Request(url, headers={"apikey": cfg["SUPABASE_KEY"],
                                               "Authorization": "Bearer " + cfg["SUPABASE_KEY"]})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read() or b"null")


def employee_full_name(cfg, emp_id):
    """ชื่อเต็มผู้ส่งรอบ — RPC คืนแค่ nickname จึงอ่านตาราง employees ด้วย service key"""
    if not emp_id:
        return None
    try:
        rows = _rest_get(cfg, "/rest/v1/employees?select=full_name&id=eq.%d" % int(emp_id))
        return F.norm((rows or [{}])[0].get("full_name")) or None
    except (urllib.error.URLError, OSError, ValueError, IndexError, TypeError) as e:
        S.log("CASH_EXCEL: อ่านชื่อพนักงาน id %s ไม่ได้ (%s) — ใช้ชื่อเล่นจาก RPC แทน"
              % (emp_id, type(e).__name__))
        return None


# =====================================================================================
# แผนการเติม (pure — ไม่แตะไฟล์ ไม่แตะเครือข่าย → เทสต์ได้)
# =====================================================================================
def _money_of(doc, express_amount):
    """แยกเงินของใบ RE/AI ตามช่องทางในแอป → (cash, noncash, noncash_channel) หรือ None ถ้ายังไม่เลือก"""
    if not doc or not doc.get("channel"):
        return None
    ch = doc["channel"]
    amt = doc.get("amount_now")
    amt = _num(amt if amt is not None else (doc.get("amount") if doc.get("amount") is not None
                                            else express_amount))
    nc = doc.get("noncash_channel") or ("qr" if ch == "qr" else "transfer")
    if ch == "cash":
        return amt, 0.0, nc
    if ch in ("transfer", "qr"):
        return 0.0, amt, ch
    if ch == "mixed":
        cash = _num(doc.get("cash_amount"))
        return cash, _num(amt - cash), nc
    return None


class Plan(object):
    def __init__(self):
        self.cells = {}       # coord -> value
        self.comments = {}    # coord -> text
        self.notes = []       # ข้อสังเกตลง log/รายงาน (ไม่มีชื่อลูกค้า)
        self.stats = {}

    def set(self, coord, value):
        if value is None or value == "":
            return
        if isinstance(value, float):
            value = round(value, 2)
            if value == int(value):
                value = int(value)
        self.cells[coord] = value

    def comment(self, coord, text):
        self.comments[coord] = (self.comments[coord] + "\n" + text) if coord in self.comments else text


def _alloc(lines, money):
    """แบ่งเงินของ RE ลงทีละรายการ (ตามลำดับ ARRCPIT) — เงินสดก่อน แล้วค่อยโอน/QR
    คืน [(line, cash, noncash, short)] · short = ส่วนที่ไม่มีเงินจริงรองรับ (RE หัก SR/AI ในตัว)"""
    cash, noncash = money[0], money[1]
    out = []
    for ln in lines:
        need = ln["amount"]
        c = min(cash, need) if need > 0 else 0.0
        cash = _num(cash - c)
        n = min(noncash, _num(need - c)) if need > 0 else 0.0
        noncash = _num(noncash - n)
        short = _num(need - c - n)
        if short > 0.005:
            # ใบเสร็จหัก SR/AI ในตัว: แบบที่พนักงานลง = IV เต็มยอดในช่องเดียวกับใบเสร็จ
            # แล้วแถว SR ลงช่องเดียวกัน (สูตร C113 หักคืน · เงินสดไม่เพี้ยนถ้า RE ไม่ใช่เงินสด)
            if money[1] > 0.005:
                n = _num(n + short)
            else:
                c = _num(c + short)
        out.append((ln, _num(c), _num(n), short))
    return out


def build_plan(ex, detail, day, full_name=None):
    """ex = load_express(), detail = cash_round_detail() (None/{} ได้) → Plan"""
    P = Plan()
    detail = detail or {}
    rnd = detail.get("round") or {}
    docs = {d.get("doc_no"): d for d in (detail.get("docs") or []) if d.get("doc_no")}
    has_round = bool(rnd)
    if not has_round:
        P.notes.append("ไม่มีรอบในแอป — เติมจาก Express อย่างเดียว ช่องทางชำระเว้นว่าง")

    P.set("F1", datetime.datetime(day.year, day.month, day.day))
    P.set("E1", TH_WEEKDAY[day.weekday()])

    # RE ที่เกี่ยวกับวันนี้: RE ใน Express ลงวันนี้ ∪ RE ในรอบของแอป (รวมใบยกมา)
    re_nos = set(ex["res"]) | {k for k, d in docs.items() if d.get("doc_type") == "RE"}

    def re_money(re_no):
        r = ex["res"].get(re_no) or {}
        return _money_of(docs.get(re_no), _num(r.get("NETAMT")))

    # จับคู่ RE → IV
    same_day_pay = {}     # iv_no -> [(re_no, cash, noncash, nc_channel|None, short)]
    paid_rows = []        # (re_no, iv_no|None, line_amount, cash, noncash, nc, short, carried)
    day_iv = {(r.get("DOCNUM") or "").strip() for r in ex["ivs"]}
    sr_col = {}           # sr_no -> (คอลัมน์|None, re_no)
    for re_no in sorted(re_nos):
        rdoc = docs.get(re_no) or {}
        r = ex["res"].get(re_no) or {}
        re_date = r.get("DOCDAT") or (datetime.date.fromisoformat(rdoc["doc_date"])
                                      if rdoc.get("doc_date") else day)
        carried = bool(rdoc.get("carried")) or re_date < day
        money = re_money(re_no)
        lines = [ln for ln in ex["lines"].get(re_no, []) if ln["rectyp"] == "3"]
        if not lines:
            amt = _num(r.get("NETAMT") if r else (rdoc.get("amount_now") or 0))
            if amt <= 0:
                continue
            m = money or (None, None, None)
            paid_rows.append((re_no, None, amt, m[0], m[1], m[2], 0.0, carried))
            P.notes.append("%s ไม่พบรายการใน ARRCPIT — ลงแถววางบิลด้วยเลข RE" % re_no)
            continue
        alloc = _alloc(lines, money) if money else [(ln, None, None, 0.0) for ln in lines]
        for ln in ex["lines"].get(re_no, []):
            if ln["rectyp"] == "5":      # SR ที่ถูกหักในใบเสร็จนี้ → ลงช่องเดียวกับใบเสร็จ
                sr_col[ln["doc_no"]] = (None if not money else
                                        CHANNEL_COL[money[2]] if money[1] > 0.005 else "F", re_no)
        for ln, c, n, short in alloc:
            nc = money[2] if money else None
            if ln["doc_no"] in day_iv and not carried:
                same_day_pay.setdefault(ln["doc_no"], []).append((re_no, c, n, nc, short))
            else:
                paid_rows.append((re_no, ln["doc_no"], ln["amount"], c, n, nc, short, carried))

    # ---- แถว 3–62: IV ของวัน ----
    if len(ex["ivs"]) > len(IV_ROWS):
        P.notes.append("IV %d ใบ เกิน %d แถวของฟอร์ม — ตัดท้าย" % (len(ex["ivs"]), len(IV_ROWS)))
    for row, r in zip(IV_ROWS, ex["ivs"]):
        iv = (r.get("DOCNUM") or "").strip()
        P.set("A%d" % row, iv)
        P.set("B%d" % row, cus_name(r, ex))
        P.set("C%d" % row, area_label(r, ex))
        adv = (r.get("ADVNUM") or "").strip()
        if len(adv) > 2 and _num(r.get("ADVAMT")) > 0:   # 'AI' เปล่า ๆ = ไม่มีมัดจำ
            P.set("D%d" % row, adv)
            P.set("E%d" % row, _num(r.get("ADVAMT")))
        net = _num(r.get("NETAMT"))
        sums = {"F": 0.0, "H": 0.0, "I": 0.0}
        covered, banks, unknown = 0.0, [], []
        for re_no, c, n, nc, short in same_day_pay.get(iv, []):
            if c is None:
                unknown.append(re_no)
                continue
            sums["F"] += c
            if n:
                col = CHANNEL_COL[nc]
                sums[col] += n
                if BANK_LABEL[nc] not in banks:
                    banks.append(BANK_LABEL[nc])
            covered += c + n
            if short > 0.005:
                P.comment("A%d" % row, "%s หัก SR/AI ในใบเสร็จ %s บาท" % (re_no, short))
        for col, v in sums.items():
            if v > 0.005:
                P.set("%s%d" % (col, row), v)
        if unknown:
            P.comment("F%d" % row, "%s: %s" % (NOTE_NO_CHANNEL, ", ".join(unknown)))
        elif not same_day_pay.get(iv) and net > 0.005:
            P.set("G%d" % row, net)                        # ไม่มี RE วันนี้ = เงินเชื่อ
        elif covered < net - 0.005:
            P.set("G%d" % row, _num(net - covered))        # ชำระบางส่วน ที่เหลือ = เชื่อ
        if banks:
            P.set("L%d" % row, "/".join(banks))

    # ---- แถว 65–72: AI (มัดจำ) — AI ลงวันนี้ใน Express ∪ AI ในรอบของแอป ----
    ai_rows = list(ex["ais"])
    ai_seen = {(r.get("DOCNUM") or "").strip() for r in ai_rows}
    for k, d in sorted(docs.items()):
        if d.get("doc_type") == "AI" and k not in ai_seen:
            ai_rows.append({"DOCNUM": k, "CUSCOD": d.get("customer_code") or "",
                            "NETAMT": d.get("amount_now"), "_app_only": True,
                            "_name": d.get("customer_name"), "_date": d.get("doc_date")})
    if len(ai_rows) > len(AI_ROWS):
        P.notes.append("AI %d ใบ เกิน %d แถว — ตัดท้าย" % (len(ai_rows), len(AI_ROWS)))
    for row, r in zip(AI_ROWS, ai_rows):
        no = (r.get("DOCNUM") or "").strip()
        P.set("A%d" % row, no)
        P.set("B%d" % row, cus_name(r, ex) if not r.get("_app_only") else
              (r.get("_name") or cus_name(r, ex)))
        P.set("C%d" % row, area_label(r, ex))
        money = _money_of(docs.get(no), _num(r.get("NETAMT")))
        if money is None:
            P.comment("F%d" % row, NOTE_NO_CHANNEL)
        else:
            if money[0] > 0.005:
                P.set("F%d" % row, money[0])
            if money[1] > 0.005:
                P.set("%s%d" % (CHANNEL_COL[money[2]], row), money[1])
                P.set("L%d" % row, BANK_LABEL[money[2]])
        if r.get("_app_only") or (docs.get(no) or {}).get("carried"):
            P.comment("A%d" % row, "ยกมาจาก %s" % (r.get("_date") or (docs.get(no) or {}).get("doc_date")))

    # ---- แถว 75–81: SR (คืนสินค้า/ลดหนี้) ----
    if len(ex["srs"]) > len(SR_ROWS):
        P.notes.append("SR %d ใบ เกิน %d แถว — ตัดท้าย" % (len(ex["srs"]), len(SR_ROWS)))
    for row, r in zip(SR_ROWS, ex["srs"]):
        P.set("A%d" % row, (r.get("DOCNUM") or "").strip())
        P.set("B%d" % row, cus_name(r, ex))
        m = _IV_RE.search(r.get("YOUREF") or "") or _IV_RE.search(r.get("SONUM") or "")
        if m:
            P.set("C%d" % row, m.group(0))
        col, by_re = sr_col.get((r.get("DOCNUM") or "").strip(), ("F", None))
        if col is None:
            P.comment("F%d" % row, "%s: %s (หัก SR ใบนี้)" % (NOTE_NO_CHANNEL, by_re))
        else:
            P.set("%s%d" % (col, row), _num(r.get("NETAMT")))
            if by_re:
                P.comment("A%d" % row, "หักในใบเสร็จ %s" % by_re)

    # ---- แถว 90–109: วางบิลชำระแล้ว (RE วันนี้ที่ชำระ IV วันก่อน + RE ยกมา) ----
    if len(paid_rows) > len(PAID_ROWS):
        P.notes.append("รายการวางบิล %d เกิน %d แถว — ตัดท้าย" % (len(paid_rows), len(PAID_ROWS)))
    for row, (re_no, iv_no, amt, c, n, nc, short, carried) in zip(PAID_ROWS, paid_rows):
        ivr = ex["iv_index"].get(iv_no) if iv_no else None
        src_r = ivr or ex["res"].get(re_no) or {"CUSCOD": (docs.get(re_no) or {}).get("customer_code") or ""}
        P.set("A%d" % row, iv_no or re_no)
        P.set("B%d" % row, cus_name(src_r, ex) or (docs.get(re_no) or {}).get("customer_name"))
        P.set("C%d" % row, area_label(src_r, ex))
        P.set("H%d" % row, amt)
        if c is None:
            P.comment("J%d" % row, "%s: %s" % (NOTE_NO_CHANNEL, re_no))
        else:
            if c > 0.005:
                P.set("J%d" % row, c)
            if n > 0.005:
                P.set("D%d" % row, n)
                P.set("E%d" % row, BANK_LABEL[nc])
                P.set("G%d" % row, datetime.datetime(day.year, day.month, day.day))
        if short > 0.005:
            P.comment("A%d" % row, "%s หัก SR/AI ในใบเสร็จ %s บาท" % (re_no, short))
        if carried and (docs.get(re_no) or {}).get("carried"):
            P.comment("A%d" % row, "%s ยกมาจากรอบก่อน" % re_no)

    # ---- ค่าใช้จ่าย / เงินสด / ผู้จัดทำ ----
    exps = [e for e in (detail.get("expenses") or []) if not e.get("deleted_at")]
    if len(exps) > len(EXP_ROWS):
        rest = exps[len(EXP_ROWS) - 1:]
        exps = exps[:len(EXP_ROWS) - 1] + [{"item": "อื่น ๆ รวม %d รายการ" % len(rest),
                                            "amount": sum(_num(e.get("amount")) for e in rest)}]
    for row, e in zip(EXP_ROWS, exps):
        P.set("G%d" % row, F.norm(e.get("item")))
        P.set("K%d" % row, _num(e.get("amount")))
    if has_round:
        flt = rnd.get("float_amount")
        flt = DEFAULT_FLOAT if flt is None else _num(flt)
        P.set("C123", flt)
        P.set("C130", flt)
        if rnd.get("counted_total") is not None:
            P.set("C125", _num(rnd["counted_total"]))
        P.set("C131", full_name or detail.get("submitted_by_name"))

    # ---- KPI traffic (sql/005) ----
    tr = detail.get("traffic") or rnd.get("traffic") or {}
    for k, coord in KPI_CELLS.items():
        if tr.get(k) is not None:
            P.set(coord, int(tr[k]))
    if tr.get("lost_other_note"):
        P.set(KPI_NOTE_CELL, F.norm(tr["lost_other_note"]))

    P.stats = {"iv": min(len(ex["ivs"]), len(IV_ROWS)), "ai": min(len(ai_rows), len(AI_ROWS)),
               "sr": min(len(ex["srs"]), len(SR_ROWS)), "paid": min(len(paid_rows), len(PAID_ROWS)),
               "expenses": len(exps), "no_channel": sum(1 for t in P.comments.values()
                                                       if NOTE_NO_CHANNEL in t),
               "round": rnd.get("status") if has_round else None}
    return P


# =====================================================================================
# เขียนลงสมุดงาน
# =====================================================================================
def _is_formula(cell):
    v = cell.value
    return isinstance(v, str) and v.startswith("=")


def _anchor(ws, coord):
    """เซลล์ในช่วง merge → คืนเซลล์มุมซ้ายบน (เขียนได้) · ไม่ใช่มุม = None"""
    for rng in ws.merged_cells.ranges:
        if coord in rng:
            tl = "%s%d" % (openpyxl.utils.get_column_letter(rng.min_col), rng.min_row)
            return tl if tl == coord else None
    return coord


def apply_plan(wb, day, plan):
    """สร้าง/แทนที่ชีต '<วัน>_auto' แล้วเติมตามแผน · คืน (ชื่อชีต, ชื่อชีตต้นแบบ, ข้ามสูตร)"""
    name = "%d%s" % (day.day, AUTO_SUFFIX)
    if name in wb.sheetnames:
        del wb[name]
    tpl, same_day = template_sheet(wb, day.day)
    if not tpl:
        raise RuntimeError("ไม่มีชีตวันให้ใช้เป็นต้นแบบ (%s)" % wb.sheetnames)
    src = wb[tpl]
    ws = wb.copy_worksheet(src)
    ws.title = name
    # copy_worksheet ไม่ลอก print area / freeze / tab color ครบ — ลอกที่สำคัญเอง
    ws.freeze_panes = src.freeze_panes
    ws.sheet_properties.tabColor = "FFC000"
    # ย้ายไปไว้ถัดจากชีตต้นแบบ อ่านง่ายกว่าไปต่อท้ายสุด
    wb.move_sheet(ws, offset=wb.sheetnames.index(tpl) + 1 - wb.sheetnames.index(name))

    blocked = []
    for coord in clear_targets():
        a = _anchor(ws, coord)
        if a is None:
            continue
        cell = ws[a]
        if _is_formula(cell):
            continue
        cell.value = None
        cell.comment = None
    for coord, value in plan.cells.items():
        a = _anchor(ws, coord)
        if a is None or _is_formula(ws[a]):
            blocked.append(coord)            # ไม่เขียนทับสูตร/กลาง merge เด็ดขาด
            continue
        ws[a].value = value
        if isinstance(value, datetime.datetime) and coord != "F1":
            ws[a].number_format = src["F1"].number_format or "d/mm/yyyy"
    for coord, text in plan.comments.items():
        a = _anchor(ws, coord) or coord
        c = Comment(text, COMMENT_AUTHOR)
        c.width, c.height = 240, 60
        ws[a].comment = c
    if not same_day:
        plan.notes.append("ยังไม่มีชีตวัน %d — ใช้ชีต '%s' เป็นต้นแบบ" % (day.day, tpl))
    return name, tpl, blocked


def _rotate_backups(path):
    baks = sorted(glob.glob(glob.escape(path) + ".bak_*"), key=os.path.getmtime)
    for old in baks[:-BACKUP_KEEP]:
        try:
            os.remove(old)
        except OSError:
            pass


def write_workbook(src_path, day, plan, out=None):
    """โหลด → เติม → save (atomic) → ตรวจซ้ำ · คืน dict ผล · ไม่ได้เขียน = {'skipped': เหตุผล}"""
    if not os.path.exists(src_path):
        return {"skipped": "ไม่พบไฟล์ %s" % os.path.basename(src_path)}
    lock = excel_locked(src_path)
    if lock and not out:
        return {"skipped": "ไฟล์ถูกเปิดค้าง (%s)" % lock}
    mtime0 = os.path.getmtime(src_path)
    backup = None
    if not out:
        backup = "%s.bak_%s" % (src_path, datetime.datetime.now(F.TH).strftime("%y%m%d_%H%M"))
        shutil.copy2(src_path, backup)
        _rotate_backups(src_path)
    try:
        wb = openpyxl.load_workbook(src_path)
    except PermissionError as e:
        return {"skipped": "เปิดไฟล์ไม่ได้ (%s)" % e}
    name = "%d%s" % (day.day, AUTO_SUFFIX)
    n_before = len(wb.sheetnames) - (1 if name in wb.sheetnames else 0)
    sheet, tpl, blocked = apply_plan(wb, day, plan)

    dest = out or src_path
    d = os.path.dirname(os.path.abspath(dest))
    tmp = os.path.join(d, ".%s.cashxl_tmp.xlsx" % os.path.basename(dest))
    try:
        wb.save(tmp)
        if not out:
            if os.path.getmtime(src_path) != mtime0 or excel_locked(src_path):
                raise _Skip("ไฟล์ถูกแก้/เปิดระหว่างเติม — ยกเลิก ไม่เขียนทับ")
        os.replace(tmp, dest)
    except _Skip as e:
        return {"skipped": str(e), "backup": backup}
    except PermissionError as e:
        return {"skipped": "เขียนไม่ได้ (%s)" % e, "backup": backup}
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)

    # ตรวจซ้ำ
    chk = openpyxl.load_workbook(dest)
    if len(chk.sheetnames) != n_before + 1 or sheet not in chk.sheetnames:
        raise RuntimeError("ตรวจซ้ำไม่ผ่าน: ชีต %d → %d" % (n_before, len(chk.sheetnames)))
    ws = chk[sheet]
    for coord in list(plan.cells)[:5] + ["J3", "C124", "K130"]:
        a = _anchor(ws, coord)
        if a is None:
            continue
        want = plan.cells.get(coord)
        got = ws[a].value
        if coord in plan.cells and coord not in blocked and not _same(got, want):
            raise RuntimeError("ตรวจซ้ำไม่ผ่าน: %s ได้ %r ต้องการ %r" % (coord, got, want))
    for f in ("J3", "C124", "K130"):
        if not _is_formula(ws[f]):
            raise RuntimeError("ตรวจซ้ำไม่ผ่าน: สูตร %s หาย" % f)
    return {"sheet": sheet, "template": tpl, "path": dest, "backup": backup, "blocked": blocked,
            "sheets": len(chk.sheetnames)}


class _Skip(Exception):
    pass


def _same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) < 0.005
    return a == b


# =====================================================================================
# entry points
# =====================================================================================
def run(branch, day, cfg, mode="dry", out=None, xlsx=None, show=False, detail=None):
    """เติม 1 สาขา/วัน · mode 'dry' | 'write' · คืน dict ผล (เทสต์ส่ง detail เองได้)"""
    t0 = time.time()
    src = cfg.get("SRC", "")
    if cfg.get("PROXY_URL"):
        S.PROXY.update(url=cfg["PROXY_URL"], token=cfg.get("PROXY_TOKEN", ""), branch=branch)
    else:
        S.PROXY.clear()
    if detail is None:
        detail = fetch_round(cfg, branch, day) or {}
    extra = [d["doc_no"] for d in (detail.get("docs") or []) if d.get("doc_type") == "RE"]
    F.clear_cache()
    ex = load_express(src, day, extra)
    rnd = detail.get("round") or {}
    full = employee_full_name(cfg, rnd.get("submitted_by")) if rnd.get("submitted_by") else None
    plan = build_plan(ex, detail, day, full)
    path = xlsx or excel_path(branch, day)
    res = {"branch": branch, "date": day.isoformat(), "path": path, "stats": plan.stats,
           "notes": plan.notes, "cells": len(plan.cells), "comments": len(plan.comments)}
    if show:
        for k in sorted(plan.cells, key=lambda c: (int(re.sub(r"\D", "", c)), c)):
            print("  %-5s %r%s" % (k, plan.cells[k], "  # " + plan.comments[k] if k in plan.comments else ""))
        for k in plan.comments:
            if k not in plan.cells:
                print("  %-5s (ว่าง)  # %s" % (k, plan.comments[k]))
    st = plan.stats
    summary = ("%s %s · IV %d AI %d SR %d วางบิล %d ค่าใช้จ่าย %d · ยังไม่เลือกช่องทาง %d · รอบ %s"
               % (branch, day.isoformat(), st["iv"], st["ai"], st["sr"], st["paid"], st["expenses"],
                  st["no_channel"], st["round"] or "-"))
    if mode != "write" and not out:
        S.log("CASH_EXCEL: %s (DRY) %.1f วิ" % (summary, time.time() - t0))
        res["plan"] = plan
        return res
    w = write_workbook(path, day, plan, out=out)
    res.update(w)
    res["plan"] = plan
    if w.get("skipped"):
        S.log("CASH_EXCEL: %s ข้าม — %s" % (summary, w["skipped"]))
    else:
        S.log("CASH_EXCEL: %s -> ชีต %s (ต้นแบบ %s)%s %.1f วิ" % (
            summary, w["sheet"], w["template"], " ลง %s" % out if out else "", time.time() - t0))
        if w.get("blocked"):
            S.log("CASH_EXCEL: %s ไม่เขียนทับเซลล์สูตร/merge %s" % (branch, w["blocked"]))
    for n in plan.notes:
        S.log("CASH_EXCEL: %s %s" % (branch, n))
    return res


def fill_after_export(branch, day, cfg):
    """hook ของ cash_poller หลัง cash_mark_exported — ห้ามทำให้ poller ล้ม
    เปิดด้วย env CASH_EXCEL_ENABLED=1 (ค่าเริ่มปิด จนกว่า CTO เคาะติดตั้ง)"""
    if os.environ.get("CASH_EXCEL_ENABLED") != "1":
        return None
    try:
        if isinstance(day, str):
            day = datetime.date.fromisoformat(day)
        return run(branch, day, cfg, mode="write")
    except F.RpcMissing as e:
        S.log("CASH_EXCEL: %s ข้าม — ยังไม่มี RPC %s" % (branch, e))
    except Exception as e:
        S.log("ERROR: CASH_EXCEL %s %s %s: %s" % (branch, day, type(e).__name__, e))
    return None


def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    if len(argv) < 2 or argv[0] in ("-h", "--help"):
        raise SystemExit(__doc__)
    branch = argv[0].upper()
    day = datetime.date.fromisoformat(argv[1])

    def opt(name):
        if name in argv:
            i = argv.index(name)
            if i + 1 < len(argv):
                return argv[i + 1]
            raise SystemExit("%s ต้องมีค่า" % name)
        return None
    out, xlsx = opt("--out"), opt("--file")
    mode = "write" if "--write" in argv else "dry"
    if "--dry-run" in argv and "--write" in argv:
        raise SystemExit("เลือก --dry-run หรือ --write อย่างใดอย่างหนึ่ง")
    if out and os.path.abspath(out) == os.path.abspath(xlsx or excel_path(branch, day)):
        raise SystemExit("--out ต้องไม่ใช่ไฟล์ต้นทาง (ใช้ --write แทน)")
    import cash_poller as CP
    cfg = CP.load_branch_cfg(branch)
    if openpyxl is None:
        raise SystemExit("ต้องติดตั้ง openpyxl")
    try:
        run(branch, day, cfg, mode=mode, out=out, xlsx=xlsx, show="--show" in argv)
    except Exception as e:
        S.log("ERROR: CASH_EXCEL %s %s %s: %s" % (branch, day, type(e).__name__, e))
        raise SystemExit(1)


if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
"""เทสต์ cash_excel_fill ด้วยสำเนา Excel จริง (SKN ต.ค. 69) + DBF จริง (สำเนา) + mock RPC

รัน:  /usr/bin/python3 -m unittest tests.test_cash_excel_fill -v
      CASH_TEST_XL=<โฟลเดอร์ที่มี SKN10.xlsx>  CASH_TEST_DBF=<โฟลเดอร์ที่มี SKN/>
      (ไม่ตั้ง = ใช้ scratchpad ของเซสชันที่เขียนเทสต์ · ไม่มีไฟล์ = skip)
      CASH_TEST_REPORT=1 พิมพ์ตารางเทียบรายเซลล์ ชีต "3_auto" กับ "3" (คีย์มือ)

ไม่แตะ Supabase / Drive / ~/007so_push — ทุกไฟล์ที่เขียนอยู่ใน temp
"""
import os, sys, shutil, tempfile, datetime, unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
_SP = "/private/tmp/claude-501/-Users-cto007/8f158c1c-f2fb-4fa7-a80a-45c4ffa17bd4/scratchpad"
XL_DIR = os.environ.get("CASH_TEST_XL") or os.path.join(_SP, "xl")
DBF_ROOT = os.environ.get("CASH_TEST_DBF") or os.path.join(_SP, "dbf")
DAY = datetime.date(2026, 10, 3)

_TMP_STATE = tempfile.mkdtemp(prefix="cash_xl_state_")
os.environ["LOCALAPPDATA"] = _TMP_STATE
os.environ.pop("CASH_EXCEL_ENABLED", None)

import openpyxl                 # noqa: E402
import so_push as S             # noqa: E402
import cash_feed as F           # noqa: E402
import cash_excel_fill as X     # noqa: E402

SRC_XL = os.path.join(XL_DIR, "SKN10.xlsx")
SRC_DBF = os.path.join(DBF_ROOT, "SKN")
HAVE = os.path.exists(SRC_XL) and os.path.exists(os.path.join(SRC_DBF, "ARTRN.DBF"))

COMPARE = ([("ABCDEFGHI", r) for r in range(3, 63)] + [("ABCDEFGHI", r) for r in range(65, 73)]
           + [("ABCDEFGHI", r) for r in range(75, 82)] + [("ABCDEFGHI", r) for r in range(90, 110)]
           + [("GK", r) for r in range(114, 130)])


def channel_from_youref(youref):
    """mock การเลือกช่องทางของพนักงาน: Express YOUREF ของ RE มีคำ kshop/ksme อยู่แล้ว"""
    y = (youref or "").lower()
    if "kshop" in y or "qr" in y:
        return "qr"
    if "ksme" in y or "โอน" in y:
        return "transfer"
    return "cash"


def mock_detail(day, with_005=True):
    """cash_round_detail จำลอง: RE/AI ของวันจาก DBF + ค่าใช้จ่าย/นับเงิน/traffic ตามที่พนักงานคีย์ชีต 3
    (ค่าใช้จ่าย/traffic เป็น input ของแอป ไม่ได้มาจาก Express — ใช้ค่าชีตจริงเป็น input)"""
    docs = []
    for r in S.read_dbf(os.path.join(SRC_DBF, "ARTRN.DBF"),
                        fields={"RECTYP", "DOCNUM", "DOCDAT", "NETAMT", "YOUREF", "CUSCOD", "DOCSTAT"}):
        if r.get("DOCDAT") == day and r["RECTYP"] in ("9", "0") and r.get("DOCSTAT") != "C":
            ch = channel_from_youref(r.get("YOUREF"))
            d = {"doc_no": r["DOCNUM"], "doc_type": "RE" if r["RECTYP"] == "9" else "AI",
                 "doc_date": day.isoformat(), "customer_code": r["CUSCOD"], "customer_name": None,
                 "amount_now": r["NETAMT"], "carried": False, "channel": ch,
                 "cash_amount": r["NETAMT"] if ch == "cash" else 0}
            if not with_005 and ch == "qr":          # ก่อน sql/005 ไม่มี qr → ถือเป็น transfer
                d["channel"] = "transfer"
            if with_005:
                d["noncash_channel"] = ch if ch != "cash" else None
            docs.append(d)
    det = {"round": {"id": 1, "branch": "SKN", "round_date": day.isoformat(), "status": "submitted",
                     "float_amount": 5000, "counted_total": 3167, "submitted_by": 77},
           "submitted_by_name": "แพร", "docs": docs,
           "expenses": [{"item": i, "amount": a} for i, a in (
               ("ค่าน้ำมันเครื่องตัดหญ้า", 200), ("ค่าตำรวจ สภ.เมือง จราจร", 500), ("ค่าฟิล์มยืด", 560),
               ("ค่าน้ำแข็ง", 100), ("ค่าซ่อมรถ กข 2549", 100), ("ค่าซื้อของไหว้", 475),
               ("ค่าแบตเตอร์รี่เครื่องชั่ง", 398))] + [{"item": "ลบแล้ว", "amount": 9, "deleted_at": "x"}]}
    if with_005:
        det["traffic"] = {"visitors": 1, "closed": 1, "lost_brand": 9, "lost_price": 0, "lost_no_stock": 0,
                          "lost_survey": 0, "lost_qty": 0, "lost_delivery": 0, "lost_other": 0,
                          "lost_other_note": None}
    return det


def compare(ws_auto, ws_real):
    """เทียบรายเซลล์ (นับเฉพาะเซลล์ที่มีค่าอย่างน้อยฝั่งหนึ่ง) → (ตรง, ทั้งหมด, [ต่าง])"""
    same, total, diffs = 0, 0, []
    for cols, r in COMPARE:
        for c in cols:
            a, b = ws_auto["%s%d" % (c, r)].value, ws_real["%s%d" % (c, r)].value
            if isinstance(a, str) and a.startswith("="):
                continue
            na = a.strip() if isinstance(a, str) else a
            nb = b.strip() if isinstance(b, str) else b
            if na in (None, "") and nb in (None, ""):
                continue
            total += 1
            if X._same(na, nb):
                same += 1
            else:
                diffs.append(("%s%d" % (c, r), b, a))
    return same, total, diffs


@unittest.skipUnless(HAVE, "ไม่มีสำเนา Excel/DBF ใน scratchpad")
class ExcelFillTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="cash_xl_")
        self.xl = os.path.join(self.tmp, "รายงานขายสกลนคร6910.xlsx")
        shutil.copy2(SRC_XL, self.xl)
        self.cfg = {"BRANCH": "SKN", "SRC": SRC_DBF, "SUPABASE_URL": "http://x", "SUPABASE_KEY": "service-test"}
        self._emp = X.employee_full_name
        X.employee_full_name = lambda cfg, i: "เพชรฤมล  เหลาแตว" if i == 77 else None
        self._rpc = F.RPC

        def no_net(*a, **k):
            raise AssertionError("เทสต์ห้ามเรียกเครือข่าย")
        F.RPC = no_net
        self.outroot = tempfile.mkdtemp(prefix="cash_xl_out_")
        self._env = os.environ.get("CASH_EXCEL_OUT_ROOT")
        os.environ["CASH_EXCEL_OUT_ROOT"] = self.outroot

    def tearDown(self):
        X.employee_full_name = self._emp
        F.RPC = self._rpc
        if self._env is None:
            os.environ.pop("CASH_EXCEL_OUT_ROOT", None)
        else:
            os.environ["CASH_EXCEL_OUT_ROOT"] = self._env
        shutil.rmtree(self.tmp, ignore_errors=True)
        shutil.rmtree(self.outroot, ignore_errors=True)

    def _run(self, detail, out=None, mode="write"):
        return X.run("SKN", DAY, self.cfg, mode=mode, out=out, xlsx=self.xl, detail=detail)

    def test_excel_path(self):
        p = X.excel_path("BK", datetime.date(2026, 10, 6), root="/r")
        self.assertEqual(p, "/r/รายงานขายสาขาบึงกาฬ/รายงานขายBK/รายงานขายปีBK2569/รายงานขายบึงกาฬ6910.xlsx")
        self.assertTrue(X.excel_path("SKN", datetime.date(2027, 1, 2), root="/r").endswith(
            "รายงานขายSKNปี70/รายงานขายสกลนคร7001.xlsx"))
        self.assertTrue(X.excel_path("PPS", DAY, root="/r").endswith("รายงานขายPPSปี69/รายงานขายโพนพิสัย6910.xlsx"))

    def test_day_sheet_map(self):
        wb = openpyxl.Workbook()
        for n in ("1", "4(หยุด)", "5 หยุด ", "6หยุด", "7_auto", "สรุป"):
            wb.create_sheet(n)
        m = X.day_sheet_map(wb)
        self.assertEqual(m, {1: "1", 4: "4(หยุด)", 5: "5 หยุด ", 6: "6หยุด"})

    def test_dry_run_writes_nothing(self):
        before = os.path.getmtime(self.xl)
        res = self._run(mock_detail(DAY), mode="dry")
        self.assertEqual(os.path.getmtime(self.xl), before)
        self.assertEqual(os.listdir(self.tmp), [os.path.basename(self.xl)])
        self.assertGreater(res["cells"], 50)

    def test_out_copy_and_compare_sheet3(self):
        out = os.path.join(self.outroot, "out.xlsx")
        res = self._run(mock_detail(DAY), out=out)
        self.assertEqual(res["sheet"], "3_auto")
        self.assertEqual(res["template"], "3")
        self.assertNotIn("backup", res)
        # ต้นทางไม่ถูกแตะ
        src = openpyxl.load_workbook(self.xl)
        self.assertNotIn("3_auto", src.sheetnames)
        wb = openpyxl.load_workbook(out)
        self.assertEqual(len(wb.sheetnames), len(src.sheetnames) + 1)
        self.assertEqual(wb.sheetnames.index("3_auto"), wb.sheetnames.index("3") + 1)
        auto, real = wb["3_auto"], wb["3"]
        # ชีตจริงไม่เปลี่ยนสักเซลล์
        for row_a, row_b in zip(src["3"].iter_rows(), real.iter_rows()):
            for a, b in zip(row_a, row_b):
                self.assertEqual(a.value, b.value, a.coordinate)
        # สูตรครบเหมือนต้นแบบ
        for row in real.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    self.assertEqual(auto[c.coordinate].value, c.value, c.coordinate)
        same, total, diffs = compare(auto, real)
        pct = 100.0 * same / total
        if os.environ.get("CASH_TEST_REPORT"):
            print("\nเทียบ 3_auto กับ 3: ตรง %d/%d = %.1f%%" % (same, total, pct))
            for coord, b, a in diffs:
                print("  %-5s จริง=%-28r auto=%r" % (coord, b, a))
        # ส่วนที่มาจาก Express + ช่องทาง ต้องตรงทุกใบ (เลข IV / มัดจำ / ยอดเงินแต่ละช่อง)
        for r in range(3, 20):
            for c in "ADEFGHI":
                self.assertTrue(X._same(auto["%s%d" % (c, r)].value, real["%s%d" % (c, r)].value),
                                "%s%d จริง=%r auto=%r" % (c, r, real["%s%d" % (c, r)].value,
                                                         auto["%s%d" % (c, r)].value))
        self.assertGreaterEqual(pct, 80.0)
        self.assertEqual(auto["C123"].value, 5000)
        self.assertEqual(auto["C125"].value, 3167)
        self.assertEqual(auto["C131"].value, "เพชรฤมล  เหลาแตว")
        self.assertEqual(auto["N8"].value, 9)
        self.assertEqual(auto["L4"].value, "KSHOP")
        self.assertEqual(auto["L5"].value, "KSME")
        self.assertEqual(auto["F1"].value, datetime.datetime(2026, 10, 3))

    def _sha(self, path):
        import hashlib
        with open(path, "rb") as f:
            return hashlib.sha256(f.read()).hexdigest()

    def test_write_never_touches_branch_file(self):
        """แยกไฟล์ แยกส่วนกันทำงาน (CTO 10 ต.ค. 69): ไฟล์สาขาเหมือนเดิมทุกไบต์ ไม่มีไฟล์ใหม่ในโฟลเดอร์สาขา"""
        h0, m0 = self._sha(self.xl), os.path.getmtime(self.xl)
        n0 = len(openpyxl.load_workbook(self.xl).sheetnames)
        r1 = self._run(mock_detail(DAY))
        self.assertEqual(self._sha(self.xl), h0)
        self.assertEqual(os.path.getmtime(self.xl), m0)
        self.assertEqual(os.listdir(self.tmp), [os.path.basename(self.xl)])
        want = X.auto_out_path("SKN", DAY)
        self.assertEqual(r1["path"], want)
        self.assertTrue(want.startswith(self.outroot) and os.path.exists(want))
        self.assertNotIn("backup", r1)
        r2 = self._run(mock_detail(DAY))             # รันซ้ำ = แทนที่ไฟล์ผล ไม่งอกชีต
        wb = openpyxl.load_workbook(want)
        self.assertEqual(len(wb.sheetnames), n0 + 1)
        self.assertEqual(r2["sheets"], n0 + 1)
        self.assertEqual(self._sha(self.xl), h0)
        self.assertFalse([f for f in os.listdir(os.path.dirname(want)) if "cashxl_tmp" in f])

    def test_out_inside_branch_folder_refused(self):
        with self.assertRaises(ValueError):
            self._run(mock_detail(DAY), out=os.path.join(self.tmp, "x_auto.xlsx"))
        with self.assertRaises(ValueError):
            self._run(mock_detail(DAY), out=self.xl)
        self.assertEqual(os.listdir(self.tmp), [os.path.basename(self.xl)])

    def test_old_auto_sheets_in_branch_file_not_copied(self):
        wb = openpyxl.load_workbook(self.xl)
        wb.copy_worksheet(wb[wb.sheetnames[0]]).title = "5_auto"   # จำลองชีตค้างจากเวอร์ชันเก่า
        wb.save(self.xl)
        h0 = self._sha(self.xl)
        res = self._run(mock_detail(DAY))
        names = openpyxl.load_workbook(res["path"]).sheetnames
        self.assertNotIn("5_auto", names)
        self.assertIn("%d_auto" % DAY.day, names)
        self.assertEqual(self._sha(self.xl), h0)

    def test_open_in_excel_still_reads(self):
        open(os.path.join(self.tmp, "~$" + os.path.basename(self.xl)), "w").close()
        h0 = self._sha(self.xl)
        res = self._run(mock_detail(DAY))
        self.assertTrue(os.path.exists(res["path"]))
        self.assertEqual(self._sha(self.xl), h0)

    def test_without_sql005_and_without_channel(self):
        det = mock_detail(DAY, with_005=False)
        det["docs"] = [d for d in det["docs"] if d["doc_no"] != "RE0027119"]   # ยังไม่เลือก/ไม่อยู่ในแอป
        out = os.path.join(self.outroot, "out.xlsx")
        res = self._run(det, out=out)
        ws = openpyxl.load_workbook(out)["3_auto"]
        self.assertEqual(ws["H4"].value, 13071)          # qr ไม่มี → ลงโอน
        self.assertIsNone(ws["I4"].value)
        self.assertIsNone(ws["H5"].value)                 # IV6904071 ไม่มีช่องทาง → ว่าง + comment
        self.assertIsNone(ws["G5"].value)
        self.assertIn(X.NOTE_NO_CHANNEL, ws["F5"].comment.text)
        self.assertIsNone(ws["N3"].value)                 # ไม่มี traffic
        self.assertEqual(res["stats"]["no_channel"], 1)

    def test_no_round_express_only(self):
        out = os.path.join(self.outroot, "out.xlsx")
        res = self._run({"branch": "SKN", "date": DAY.isoformat(), "round": None}, out=out)
        ws = openpyxl.load_workbook(out)["3_auto"]
        self.assertEqual(ws["A3"].value, "IV6904069")
        self.assertEqual(ws["G3"].value, 3225)
        self.assertIsNone(ws["C123"].value)
        self.assertIsNone(ws["G114"].value)
        self.assertTrue(res["notes"])

    def test_mixed_channel_split(self):
        det = mock_detail(DAY)
        for d in det["docs"]:
            if d["doc_no"] == "RE0027117":                # IV6904073 29,847
                d.update(channel="mixed", cash_amount=10000, noncash_channel="qr")
        out = os.path.join(self.outroot, "out.xlsx")
        self._run(det, out=out)
        ws = openpyxl.load_workbook(out)["3_auto"]
        self.assertEqual(ws["F7"].value, 10000)
        self.assertEqual(ws["I7"].value, 19847)
        self.assertIsNone(ws["G7"].value)

    def test_formula_cells_never_written(self):
        p = X.Plan()
        p.cells.update({"J3": 1, "C124": 2, "K90": 3, "A3": "IV1", "D3": "x", "E2": "mid-merge"})
        wb = openpyxl.load_workbook(self.xl)
        name, tpl, blocked = X.apply_plan(wb, DAY, p)
        ws = wb[name]
        self.assertEqual(sorted(blocked), ["C124", "E2", "J3", "K90"])
        self.assertEqual(ws["J3"].value, "=SUM(E3:I3)")
        self.assertEqual(ws["A3"].value, "IV1")

    def test_poller_hook_disabled_by_default(self):
        self.assertIsNone(X.fill_after_export("SKN", DAY.isoformat(), self.cfg))


if __name__ == "__main__":
    unittest.main()
